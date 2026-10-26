"""
PDF Invoice -> Excel Extractor
------------------------------
Reads a folder of invoice PDFs (different suppliers, different layouts, EN or RO)
and builds one Excel workbook with:
  * Invoices    - one row per invoice (number, date, supplier, subtotal, VAT, total, currency)
  * Line Items  - every product line from every invoice's table
  * Checks      - automatic validation: line items vs subtotal, subtotal + VAT vs total

Mismatches are highlighted in red, so wrong invoices are caught before they are paid.

Usage:
    python extract_invoices.py sample_invoices -o output/invoices.xlsx
"""
from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
import pdfplumber
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


# --- Layout profiles: add a new supplier format by adding one profile --------------
@dataclass
class Layout:
    name: str
    detect: str            # text that identifies this layout
    invoice_no: str        # regex patterns, first group = value
    date: str
    supplier: str
    subtotal: str
    vat: str
    total: str
    currency: str
    decimal_comma: bool    # True for "1.234,56" style numbers
    columns: dict          # table header -> standard field


LAYOUTS = [
    Layout(
        name="English invoice",
        detect="INVOICE",
        invoice_no=r"Invoice No:\s*(\S+)",
        date=r"Invoice Date:\s*([\d\-./]+)",
        supplier=r"^(.+)\n",
        subtotal=r"Subtotal:\s*([\d.,]+)",
        vat=r"^VAT[^:]*:\s*([\d.,]+)",
        total=r"Total Due:\s*([\d.,]+)",
        currency=r"Total Due:\s*[\d.,]+\s*([A-Z]{3})",
        decimal_comma=False,
        columns={"Description": "product", "Qty": "qty", "Unit Price": "unit_price", "Amount": "amount"},
    ),
    Layout(
        name="Romanian factura",
        detect="FACTURA",
        invoice_no=r"Factura nr\.\s*(\S+)",
        date=r"Data emiterii:\s*([\d.\-/]+)",
        supplier=r"Furnizor:\s*(.+)",
        subtotal=r"Total fara TVA:\s*([\d.,]+)",
        vat=r"^TVA[^:]*:\s*([\d.,]+)",
        total=r"Total de plata:\s*([\d.,]+)",
        currency=r"Total de plata:\s*[\d.,]+\s*([A-Z]{3})",
        decimal_comma=True,
        columns={"Denumire produs": "product", "Cant.": "qty", "Pret unitar": "unit_price", "Valoare": "amount"},
    ),
]


def to_number(text: str | None, decimal_comma: bool) -> float | None:
    if not text:
        return None
    s = text.strip()
    s = s.replace(".", "").replace(",", ".") if decimal_comma else s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def find(pattern: str, text: str) -> str | None:
    m = re.search(pattern, text, flags=re.MULTILINE)
    return m.group(1).strip() if m else None


def parse_invoice(path: Path) -> tuple[dict, list[dict]]:
    with pdfplumber.open(path) as pdf:
        text = "\n".join(p.extract_text() or "" for p in pdf.pages)
        tables = [t for p in pdf.pages for t in p.extract_tables()]

    layout = next((l for l in LAYOUTS if l.detect in text), None)
    if layout is None:
        return {"file": path.name, "status": "UNKNOWN LAYOUT"}, []

    num = lambda pat: to_number(find(pat, text), layout.decimal_comma)
    header = {
        "file": path.name,
        "layout": layout.name,
        "invoice_no": find(layout.invoice_no, text),
        "date": pd.to_datetime(find(layout.date, text), dayfirst=layout.decimal_comma, errors="coerce"),
        "supplier": find(layout.supplier, text),
        "subtotal": num(layout.subtotal),
        "vat": num(layout.vat),
        "total": num(layout.total),
        "currency": find(layout.currency, text),
    }

    items = []
    for table in tables:
        head = [c.strip() if c else "" for c in table[0]]
        if not set(layout.columns) <= set(head):
            continue
        idx = {std: head.index(col) for col, std in layout.columns.items()}
        for row in table[1:]:
            items.append({
                "invoice_no": header["invoice_no"],
                "product": row[idx["product"]],
                "qty": to_number(row[idx["qty"]], layout.decimal_comma),
                "unit_price": to_number(row[idx["unit_price"]], layout.decimal_comma),
                "amount": to_number(row[idx["amount"]], layout.decimal_comma),
                "currency": header["currency"],
            })
    return header, items


def validate(inv: pd.DataFrame, items: pd.DataFrame) -> pd.DataFrame:
    lines_sum = items.groupby("invoice_no")["amount"].sum().round(2)
    checks = inv[["invoice_no", "supplier", "subtotal", "vat", "total"]].copy()
    checks["sum_of_lines"] = checks["invoice_no"].map(lines_sum)
    checks["lines_match_subtotal"] = (checks["sum_of_lines"] - checks["subtotal"]).abs() < 0.02
    checks["subtotal_plus_vat_match_total"] = (checks["subtotal"] + checks["vat"] - checks["total"]).abs() < 0.02
    checks["status"] = checks.apply(
        lambda r: "OK" if r.lines_match_subtotal and r.subtotal_plus_vat_match_total
        else f"CHECK: total off by {r.total - r.subtotal - r.vat:+,.2f}", axis=1)
    return checks


def style(ws, money=()):
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.font = Font(bold=True, color="FFFFFF")
    for i, col in enumerate(ws.iter_cols(), 1):
        ws.column_dimensions[get_column_letter(i)].width = min(
            max(len(str(c.value)) if c.value is not None else 0 for c in col) + 3, 40)
        if col[0].value in money:
            for c in col[1:]:
                c.number_format = "#,##0.00"
        if col[0].value == "Date":
            for c in col[1:]:
                c.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    ap = argparse.ArgumentParser(description="Extract invoice PDFs into one validated Excel file.")
    ap.add_argument("folder", type=Path)
    ap.add_argument("-o", "--output", type=Path, default=Path("output/invoices.xlsx"))
    args = ap.parse_args()

    headers, all_items = [], []
    for pdf in sorted(args.folder.glob("*.pdf")):
        h, items = parse_invoice(pdf)
        headers.append(h)
        all_items.extend(items)
        print(f"  {pdf.name:<28} {h.get('layout', '?'):<18} {len(items)} line items")

    inv = pd.DataFrame(headers)
    items = pd.DataFrame(all_items)
    checks = validate(inv, items)

    inv_view = inv.rename(columns={"file": "Source PDF", "layout": "Layout", "invoice_no": "Invoice No",
                                   "date": "Date", "supplier": "Supplier", "subtotal": "Subtotal",
                                   "vat": "VAT", "total": "Total", "currency": "Currency"})
    items_view = items.rename(columns={"invoice_no": "Invoice No", "product": "Product", "qty": "Qty",
                                       "unit_price": "Unit Price", "amount": "Amount", "currency": "Currency"})
    checks_view = checks.rename(columns={"invoice_no": "Invoice No", "supplier": "Supplier", "subtotal": "Subtotal",
                                         "vat": "VAT", "total": "Total", "sum_of_lines": "Sum of Lines",
                                         "lines_match_subtotal": "Lines = Subtotal",
                                         "subtotal_plus_vat_match_total": "Subtotal + VAT = Total",
                                         "status": "Status"})

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(args.output, engine="openpyxl") as w:
        inv_view.to_excel(w, sheet_name="Invoices", index=False)
        items_view.to_excel(w, sheet_name="Line Items", index=False)
        checks_view.to_excel(w, sheet_name="Checks", index=False)
        style(w.sheets["Invoices"], money={"Subtotal", "VAT", "Total"})
        style(w.sheets["Line Items"], money={"Unit Price", "Amount"})
        style(w.sheets["Checks"], money={"Subtotal", "VAT", "Total", "Sum of Lines"})
        ws = w.sheets["Checks"]
        for row in ws.iter_rows(min_row=2):
            ok = row[-1].value == "OK"
            row[-1].fill = PatternFill("solid", fgColor="D9F2E3" if ok else "F8D7DA")
            row[-1].font = Font(bold=True, color="1F7A4D" if ok else "B42318")

    bad = (checks["status"] != "OK").sum()
    print(f"\n{len(inv)} invoices, {len(items)} line items extracted. {bad} invoice(s) need checking.")
    print(f"Saved to {args.output}")


if __name__ == "__main__":
    main()
